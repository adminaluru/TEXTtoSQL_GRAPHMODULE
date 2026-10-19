"""Excel reader: one text block per visible row (D-010).

- The value Excel last saved is used; the formula text is kept beside it.
  FlowOne never recalculates a formula.
- Hidden sheets, rows and columns are skipped and counted.
- Links to other workbooks are not followed; macros never run (openpyxl
  cannot run them, and macro files are refused earlier).
"""
import io
from datetime import date, datetime, time

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.formula import ArrayFormula

from ..errors import DocumentRejected
from ..limits import IngestionLimits
from ..models import Block, Location
from ..text import clean_text


def read_xlsx(data: bytes, limits: IngestionLimits) -> tuple[list[Block], list[str]]:
    _check_size(data, limits)
    # Two views of the same file: saved values, and formula text.
    values = load_workbook(io.BytesIO(data), data_only=True, keep_links=False)
    formulas = load_workbook(io.BytesIO(data), data_only=False, keep_links=False)

    blocks: list[Block] = []
    hidden_sheets = hidden_rows = hidden_columns = 0
    for sheet in values.worksheets:
        if sheet.sheet_state != "visible":
            hidden_sheets += 1
            continue
        formula_sheet = formulas[sheet.title]
        skipped_columns = _hidden_columns(sheet)
        hidden_columns += len(skipped_columns)
        headers: dict[int, str] | None = None

        for row in sheet.iter_rows():
            if not row:
                continue
            number = row[0].row
            if sheet.row_dimensions[number].hidden:
                if any(c.value is not None for c in row):
                    hidden_rows += 1
                continue

            cells = []
            for cell in row:
                if cell.column in skipped_columns:
                    continue
                formula = formula_sheet.cell(row=number, column=cell.column).value
                if cell.value is not None or _formula(formula):
                    cells.append((cell.column, cell.value, formula))
            if not cells:
                continue

            labels_only = all(isinstance(v, str) and not _formula(f) for _, v, f in cells)
            if headers is None and labels_only and len(cells) == 1:
                # A title line above the table.
                text = cells[0][1]
            elif headers is None and labels_only:
                # The first row of two or more text labels names the columns.
                headers = {column: clean_text(v) for column, v, _ in cells}
                text = " | ".join(headers.values())
            else:
                headers = headers or {}
                text = "; ".join(
                    f"{headers.get(column) or get_column_letter(column)}: {_render(value, formula)}"
                    for column, value, formula in cells
                )
            text = clean_text(text)
            if text:
                blocks.append(Block(text, Location("sheet", number, number, sheet.title)))

    notes = []
    if hidden_sheets:
        notes.append(f"{hidden_sheets} hidden sheet(s) were skipped.")
    if hidden_rows:
        notes.append(f"{hidden_rows} hidden row(s) were skipped.")
    if hidden_columns:
        notes.append(f"{hidden_columns} hidden column(s) were skipped.")
    return blocks, notes


def _check_size(data: bytes, limits: IngestionLimits) -> None:
    # A cheap streaming pass reads only the declared sheet sizes.
    workbook = load_workbook(io.BytesIO(data), read_only=True, keep_links=False)
    try:
        cells = 0
        for sheet in workbook.worksheets:
            if sheet.max_row and sheet.max_column:
                cells += sheet.max_row * sheet.max_column
                continue
            # No declared size: count while streaming, stop once over the limit.
            for row in sheet.iter_rows():
                cells += len(row)
                if cells > limits.max_sheet_cells:
                    break
    finally:
        workbook.close()
    if cells > limits.max_sheet_cells:
        raise DocumentRejected(
            "too_complex", f"The workbook has more than {limits.max_sheet_cells:,} cells.")


def _hidden_columns(sheet) -> set[int]:
    hidden: set[int] = set()
    for dimension in sheet.column_dimensions.values():
        if dimension.hidden and dimension.min:
            hidden.update(range(dimension.min, (dimension.max or dimension.min) + 1))
    return hidden


def _formula(value) -> str | None:
    if isinstance(value, ArrayFormula):
        return value.text
    if isinstance(value, str) and value.startswith("="):
        return value
    return None


def _render(value, formula_value) -> str:
    formula = _formula(formula_value)
    if formula and value is None:
        return f"no saved value (formula {formula})"
    shown = _plain(value)
    return f"{shown} (formula {formula})" if formula else shown


def _plain(value) -> str:
    if isinstance(value, datetime):
        return value.date().isoformat() if value.time() == time(0) else value.isoformat(sep=" ")
    if isinstance(value, (date, time)):
        return value.isoformat()
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    return str(value)
